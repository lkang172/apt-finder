import logging
import re
from dataclasses import dataclass
from datetime import date, timedelta
from pathlib import Path
from urllib.parse import urljoin

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

from aptfinder.http import FetchError, PoliteClient

log = logging.getLogger(__name__)

SOURCE_NAME = "California Department of Finance, E-1 Population Estimates for Cities, Counties, and the State"
E1_PAGE_URL = "https://dof.ca.gov/forecasting/demographics/estimates-e1/"
E1_XLSX_URL = "https://dof.ca.gov/media/docs/forecasting/Demographics/estimates-e1/E-1_2026_InternetVersion.xlsx"
E1_LINK = re.compile(r'href="([^"]*E-1_(\d{4})_InternetVersion\.xlsx)"')
HEADER_DATE = re.compile(r"(\d{1,2})/(\d{1,2})/(\d{4})")
RELEASED = re.compile(r"Released:\s*(.+)")
# E-1 lists no city rows for these; the county row is also the city.
CONSOLIDATED_CITY_COUNTIES = ("San Francisco",)


@dataclass(frozen=True)
class CityPopulation:
    county: str
    name: str
    population: int


@dataclass(frozen=True)
class PopulationEstimates:
    as_of: date
    released: str | None
    state_population: int
    cities: list[CityPopulation]


def discover_e1_url(client: PoliteClient, ttl: timedelta) -> str:
    try:
        fetch = client.get(E1_PAGE_URL, ttl=ttl)
    except FetchError as exc:
        log.warning("Could not load the DOF E-1 page (%s); using %s", exc, E1_XLSX_URL)
        return E1_XLSX_URL
    links = [(int(year), urljoin(E1_PAGE_URL, href)) for href, year in E1_LINK.findall(fetch.text)]
    return max(links)[1] if links else E1_XLSX_URL


def parse_e1_workbook(path: Path, target_year: int) -> PopulationEstimates:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        county_sheet = _sheet(workbook.worksheets, "E-1 CountyState")
        city_sheet = _sheet(workbook.worksheets, "E-1 CityCounty")
        county_rows, county_dates = _table(county_sheet)
        city_rows, city_dates = _table(city_sheet)
        if county_dates != city_dates:
            raise ValueError("E-1 county and city sheets report different estimate dates")
        column = min(range(len(city_dates)), key=lambda i: (abs(city_dates[i].year - target_year), -city_dates[i].year))
        released = _released(workbook.worksheets[0])
    finally:
        workbook.close()

    counties = dict(county_rows)
    state_values = counties.pop("California")
    cities = [
        CityPopulation(name, name, counties[name][column]) for name in CONSOLIDATED_CITY_COUNTIES if name in counties
    ]
    county: str | None = None
    for name, values in city_rows:
        if name == "California":
            continue
        if counties.get(name) == values:
            county = name
        elif county is None:
            raise ValueError(f"E-1 city row {name!r} appears before any county row")
        elif name.lower() != "unincorporated":
            cities.append(CityPopulation(county, name, values[column]))
    return PopulationEstimates(city_dates[column], released, state_values[column], cities)


def _sheet(sheets: list[Worksheet], prefix: str) -> Worksheet:
    for sheet in sheets:
        if sheet.title.startswith(prefix):
            return sheet
    raise ValueError(f"E-1 workbook has no sheet starting with {prefix!r}")


def _table(sheet: Worksheet) -> tuple[list[tuple[str, tuple[int, ...]]], list[date]]:
    rows = sheet.iter_rows(values_only=True)
    columns: list[int] = []
    dates: list[date] = []
    for row in rows:
        header_dates = [(i, HEADER_DATE.search(str(cell))) for i, cell in enumerate(row) if cell is not None]
        found = [(i, m) for i, m in header_dates if m]
        if found:
            columns = [i for i, _ in found]
            dates = [date(int(m.group(3)), int(m.group(1)), int(m.group(2))) for _, m in found]
            break
    if not columns:
        raise ValueError(f"E-1 sheet {sheet.title!r} has no population date header")

    table = []
    for row in rows:
        name = row[0].strip() if isinstance(row[0], str) else None
        values = tuple(_whole_number(row[i]) if i < len(row) else None for i in columns)
        if name and all(v is not None for v in values):
            table.append((name, values))
    return table, dates


def _whole_number(value: object) -> int | None:
    if isinstance(value, bool) or not isinstance(value, int | float) or value != int(value):
        return None
    return int(value)


def _released(about_sheet: Worksheet) -> str | None:
    for row in about_sheet.iter_rows(values_only=True):
        for cell in row:
            match = RELEASED.search(str(cell)) if isinstance(cell, str) else None
            if match:
                return match.group(1).strip()
    return None
